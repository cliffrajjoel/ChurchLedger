import io
import json
import time
import zipfile
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image, ImageOps
import streamlit as st
from google import genai
from google.genai import types

st.set_page_config(
    page_title="Church Ledger Digitizer", page_icon="⛪", layout="wide"
)

st.title("⛪ CSI St. Paul's Church - Autonomous Ledger Digitizer")
st.write(
    "Upload 20+ ledger photos or a `.zip` file. The app processes each page at a safe pace within the Free Tier, handles retries automatically, and compiles the multi-tab Excel file."
)

with st.sidebar:
    st.header("🔑 Settings")
    secret_key = (
        st.secrets.get("GEMINI_API_KEY", "") if hasattr(st, "secrets") else ""
    )
    api_key = st.text_input(
        "Enter Gemini API Key",
        value=secret_key,
        type="password",
        help="Get an API key from Google AI Studio (aistudio.google.com)",
    )
    st.markdown("---")
    st.info(
        "💡 **Free Tier Mode**: Paced at ~16 seconds per page to guarantee zero quota errors. A 20-image batch will take ~6 to 7 minutes."
    )

# ---------------------------------------------------------
# Schema Definition
# ---------------------------------------------------------
LEDGER_SCHEMA = {
    "type": "OBJECT",
    "properties": {
        "register_date": {
            "type": "STRING",
            "description": "Date of the register, e.g. 05-04-2026",
        },
        "members": {
            "type": "ARRAY",
            "items": {
                "type": "OBJECT",
                "properties": {
                    "s_no": {"type": "INTEGER"},
                    "sub_no": {
                        "type": "STRING",
                        "description": "Subscription No e.g. 5A",
                    },
                    "name": {"type": "STRING"},
                    "month": {"type": "STRING"},
                    "subscription": {"type": "NUMBER"},
                    "evangelism": {"type": "NUMBER"},
                    "renovation": {"type": "NUMBER"},
                    "rural_fund": {"type": "NUMBER"},
                    "poor_fund": {"type": "NUMBER"},
                    "bible_society": {"type": "NUMBER"},
                    "building_fund": {"type": "NUMBER"},
                    "sunday_school": {"type": "NUMBER"},
                    "youth_fellowship": {"type": "NUMBER"},
                    "womens_fellowship": {"type": "NUMBER"},
                    "mens_fellowship": {"type": "NUMBER"},
                    "hostel_children": {"type": "NUMBER"},
                    "educational_fund": {"type": "NUMBER"},
                    "choir_fund": {"type": "NUMBER"},
                    "alter_decoration": {"type": "NUMBER"},
                    "cemetery_maintenance": {"type": "NUMBER"},
                    "written_total": {
                        "type": "NUMBER",
                        "description": "Total amount written at rightmost column",
                    },
                },
                "required": ["s_no", "name", "written_total"],
            },
        },
    },
    "required": ["register_date", "members"],
}


def preprocess_image(img: Image.Image) -> bytes:
    try:
        img = ImageOps.exif_transpose(img)
    except Exception:
        pass

    if img.mode in ("RGBA", "P"):
        img = img.convert("RGB")

    max_dim = 2048
    if max(img.size) > max_dim:
        img.thumbnail((max_dim, max_dim), Image.Resampling.LANCZOS)

    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=85)
    return buf.getvalue()


def extract_ledger_from_image(
    img: Image.Image, client: genai.Client, status_placeholder, max_retries: int = 5
):
    prompt = """
    You are an expert church auditor transcribing an official Church Monthly Subscription and Other Fund Register.
    Extract every single row accurately.
    
    COLUMNS:
    Col 1: Subscription, Col 2: Evangelism, Col 3: Renovation, Col 4: Rural Fund,
    Col 5: Poor Fund, Col 6: Bible Society, Col 7: Building Fund, Col 8: Sunday School,
    Col 9: Youth Fellowship, Col 10: Women's Fellowship, Col 11: Men's Fellowship,
    Col 12: Hostel Children, Col 13: Educational Fund, Col 14: Choir Fund,
    Col 15: Alter Decoration H.C. Elements, Col 16: Cemetery Maintenance.
    
    RULES:
    1. If a cell is blank or empty, set its value to 0.
    2. If an amount is struck out with a pen, treat it as 0 (or the written correction).
    3. Capture the exact written total at the end of each row.
    4. Capture the date of the register accurately (e.g. 05-04-2026).
    """
    img_bytes = preprocess_image(img)

    for attempt in range(max_retries):
        try:
            response = client.models.generate_content(
                model="gemini-3.8-flash",
                contents=[
                    types.Part.from_bytes(data=img_bytes, mime_type="image/jpeg"),
                    prompt,
                ],
                config=types.GenerateContentConfig(
                    response_mime_type="application/json",
                    response_schema=LEDGER_SCHEMA,
                ),
            )
            if response.text:
                return json.loads(response.text)
        except Exception as e:
            err_str = str(e)
            if "429" in err_str or "RESOURCE_EXHAUSTED" in err_str or "503" in err_str:
                status_placeholder.warning(
                    f"⚠️ Quota paused by Google. Sleeping 65s before retry (Attempt {attempt+1}/{max_retries})..."
                )
                time.sleep(65)
                continue
            else:
                raise e

    raise RuntimeError(
        "Page failed after multiple retries due to persistent API quota limits."
    )


def build_workbook(extracted_tables):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    headers = [
        "S.No", "Sub No", "Name", "Month",
        "Subscription (1)", "Evangelism (2)", "Renovation (3)", "Rural Fund (4)",
        "Poor Fund (5)", "Bible Society (6)", "Building Fund (7)", "Sunday School (8)",
        "Youth Fellowship (9)", "Women's Fellowship (10)", "Men's Fellowship (11)",
        "Hostel Children (12)", "Educational Fund (13)", "Choir Fund (14)",
        "Alter Decoration H.C. Elements (15)", "Cemetery Maintenance (16)",
        "Calculated Total", "Written Total", "Audit Status",
    ]

    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    discrepancy_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    discrepancy_font = Font(name="Calibri", size=11, bold=True, color="C00000")
    total_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
    total_font = Font(name="Calibri", size=11, bold=True)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    double_border = Border(
        top=Side(style="thin", color="000000"),
        bottom=Side(style="double", color="000000"),
    )

    for table_data in extracted_tables:
        raw_date = str(table_data.get("register_date") or "Register").strip()
        tab_name = raw_date.replace("/", "-")
        for ch in [":", "*", "?", "[", "]", "\\", "/"]:
            tab_name = tab_name.replace(ch, "")
        tab_name = tab_name[:28] or "Register"

        orig_name = tab_name
        counter = 1
        while tab_name in wb.sheetnames:
            tab_name = f"{orig_name[:24]}_{counter}"
            counter += 1

        ws = wb.create_sheet(title=tab_name)
        ws.append(headers)

        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=1, column=col_idx)
            c.fill = header_fill
            c.font = header_font
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        members = table_data.get("members", [])
        for r_idx, m in enumerate(members, start=2):
            calc_formula = f"=SUM(E{r_idx}:T{r_idx})"
            written = float(m.get("written_total") or 0.0)

            row_vals = [
                m.get("s_no", r_idx - 1),
                m.get("sub_no", ""),
                m.get("name", ""),
                m.get("month", ""),
                float(m.get("subscription") or 0.0),
                float(m.get("evangelism") or 0.0),
                float(m.get("renovation") or 0.0),
                float(m.get("rural_fund") or 0.0),
                float(m.get("poor_fund") or 0.0),
                float(m.get("bible_society") or 0.0),
                float(m.get("building_fund") or 0.0),
                float(m.get("sunday_school") or 0.0),
                float(m.get("youth_fellowship") or 0.0),
                float(m.get("womens_fellowship") or 0.0),
                float(m.get("mens_fellowship") or 0.0),
                float(m.get("hostel_children") or 0.0),
                float(m.get("educational_fund") or 0.0),
                float(m.get("choir_fund") or 0.0),
                float(m.get("alter_decoration") or 0.0),
                float(m.get("cemetery_maintenance") or 0.0),
                calc_formula,
                written,
                f'=IF(U{r_idx}=V{r_idx}, "Balanced", "Variance")',
            ]
            ws.append(row_vals)

            item_sum = sum(float(x or 0.0) for x in row_vals[4:20])

            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=r_idx, column=col_idx)
                cell.border = thin_border
                if 5 <= col_idx <= 22:
                    cell.number_format = "#,##0"

            if abs(item_sum - written) > 0.01:
                ws.cell(row=r_idx, column=21).fill = discrepancy_fill
                ws.cell(row=r_idx, column=22).fill = discrepancy_fill
                status_cell = ws.cell(row=r_idx, column=23)
                status_cell.fill = discrepancy_fill
                status_cell.font = discrepancy_font

        num_m = len(members)
        tot_row_idx = num_m + 2
        tot_row = ["Total", "", "", ""]
        cols_letters = [
            "E", "F", "G", "H", "I", "J", "K", "L", "M", "N",
            "O", "P", "Q", "R", "S", "T", "U", "V"
        ]
        for letter in cols_letters:
            tot_row.append(f"=SUM({letter}2:{letter}{num_m+1})")
        tot_row.append("")
        ws.append(tot_row)

        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=tot_row_idx, column=col_idx)
            c.fill = total_fill
            c.font = total_font
            c.border = double_border
            if 5 <= col_idx <= 22:
                c.number_format = "#,##0"

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


# ---------------------------------------------------------
# UI Handling & State Persistence
# ---------------------------------------------------------
if "excel_data" not in st.session_state:
    st.session_state.excel_data = None
if "processed_pages_count" not in st.session_state:
    st.session_state.processed_pages_count = 0

uploaded_files = st.file_uploader(
    "Drop Ledger Images (.jpg, .png) or a .zip file here",
    type=["jpg", "jpeg", "png", "zip"],
    accept_multiple_files=True,
)

if uploaded_files:
    images_to_process = []
    for f in uploaded_files:
        if f.name.lower().endswith(".zip"):
            with zipfile.ZipFile(f) as z:
                for z_name in sorted(z.namelist()):
                    if z_name.lower().endswith((".jpg", ".jpeg", ".png")) and not z_name.startswith("__MACOSX"):
                        img_data = z.read(z_name)
                        images_to_process.append((z_name, Image.open(io.BytesIO(img_data))))
        else:
            images_to_process.append((f.name, Image.open(f)))

    total_images = len(images_to_process)
    estimated_mins = round((total_images * 16) / 60, 1)
    st.success(f"📁 Loaded **{total_images} ledger pages**. Estimated Free Tier runtime: ~**{estimated_mins} minutes**.")

    if st.button("🚀 Start Unattended Processing"):
        if not api_key:
            st.error("Please enter your Gemini API Key in the sidebar.")
        else:
            client = genai.Client(api_key=api_key)
            progress_bar = st.progress(0)
            status_text = st.empty()
            audit_log = st.container()
            extracted_results = []

            for idx, (filename, img) in enumerate(images_to_process):
                status_text.info(f"⏳ Processing page {idx+1} of {total_images}: **{filename}**...")
                try:
                    res = extract_ledger_from_image(img, client, status_text)
                    if res and res.get("members"):
                        extracted_results.append(res)
                        audit_log.write(
                            f"✔️ **Page {idx+1} ({filename})**: Date `{res.get('register_date')}` with {len(res.get('members'))} rows."
                        )
                except Exception as e:
                    audit_log.error(f"❌ Failed on {filename}: {str(e)}")

                progress_bar.progress((idx + 1) / total_images)

                # Safe pacing: 16 seconds ensures strictly < 4 calls per minute
                if idx < total_images - 1:
                    status_text.caption(f"Sleeping 16s for Free Tier quota safety before next page...")
                    time.sleep(16)

            if extracted_results:
                status_text.success("🎉 All pages completed! Building final Excel file...")
                st.session_state.excel_data = build_workbook(extracted_results).getvalue()
                st.session_state.processed_pages_count = len(extracted_results)
            else:
                st.error("⚠️ No pages could be extracted. Please check the logs above.")

# Show persistent download button if results are ready
if st.session_state.excel_data:
    st.markdown("---")
    st.success(f"✅ Ready! **{st.session_state.processed_pages_count} pages** audited and converted.")
    st.download_button(
        label="📥 Download Completed Church Register (.xlsx)",
        data=st.session_state.excel_data,
        file_name="CSI_St_Pauls_Church_Register_Audited.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
